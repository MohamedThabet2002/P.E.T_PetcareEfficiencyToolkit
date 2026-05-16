"""
Excel (.xlsx) export helpers for the PET Application.
Requires the openpyxl package (used only when a workbook is actually generated).
"""

import logging

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    HAS_OPENPYXL = True
except ImportError:  # pragma: no cover - exercised only without the dependency
    HAS_OPENPYXL = False

if HAS_OPENPYXL:
    _HEADER_FILL = PatternFill("solid", fgColor="2F6FBF")
    _HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    _TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1F2A44")
    _SECTION_FONT = Font(name="Segoe UI", size=11, bold=True, color="1F2A44")
    _CELL_FONT = Font(name="Segoe UI", size=10, color="333333")
    _META_FONT = Font(name="Segoe UI", size=9, color="666666")
    _THIN = Side(style="thin", color="DDDDDD")
    _CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
    _ACCENT = PatternFill("solid", fgColor="EAF1FB")


def _auto_width(ws, col, rows, extra=2):
    """Approximates an ideal column width from the widest cell value."""
    largest = 10
    for row in rows:
        if col < len(row) and row[col] is not None:
            width = len(str(row[col])) + extra
            if width > largest:
                largest = width
        if largest > 34:
            largest = 34
            break
    ws.column_dimensions[get_column_letter(col + 1)].width = largest


def write_report_workbook(file_path, report_title, subtitle, sections):
    """Writes a styled, professional-looking report workbook.

    Args:
        file_path: Destination path (must end in .xlsx).
        report_title: Overall report heading text.
        subtitle: Detail line under the title (dates, filters, generated-by).
        sections: List of dicts:
            {"title": str, "headers": [str, ...],
             "rows": [[cell, ...], ...],
             "summary": [(label, value), ...]}

    Returns:
        True on success, False on failure (logs the error).
    """
    if not HAS_OPENPYXL:
        logger.error("openpyxl is not installed; cannot export Excel workbook")
        return False

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Report"

        # Base row (title) is styled via cell writes; we compose a clean grid.
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        title_cell = ws.cell(row=1, column=1, value=report_title)
        title_cell.font = _TITLE_FONT
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
        sub_cell = ws.cell(row=2, column=1, value=subtitle)
        sub_cell.font = _META_FONT
        ws.row_dimensions[2].height = 14

        row_index = 4
        for section in sections:
            headers = list(section.get("headers", []))
            rows = list(section.get("rows", []))
            summary = list(section.get("summary", []))

            ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=max(len(headers), 2))
            section_cell = ws.cell(row=row_index, column=1, value=section.get("title", ""))
            section_cell.font = _SECTION_FONT
            section_cell.fill = _ACCENT
            section_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row_index].height = 18
            row_index += 1

            if headers:
                header_row = row_index
                for col, header in enumerate(headers):
                    cell = ws.cell(row=header_row, column=col + 1, value=header)
                    cell.font = _HEADER_FONT
                    cell.fill = _HEADER_FILL
                    cell.border = _CELL_BORDER
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                row_index += 1

            for data_row in rows:
                for col, value in enumerate(data_row):
                    cell = ws.cell(row=row_index, column=col + 1, value=value)
                    cell.font = _CELL_FONT
                    cell.border = _CELL_BORDER
                    if col == 0:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                row_index += 1

            for label, value in summary:
                ws.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=max(len(headers), 2) - 1)
                label_cell = ws.cell(row=row_index, column=1, value=label)
                label_cell.font = _SECTION_FONT
                value_cell = ws.cell(row=row_index, column=max(len(headers), 2), value=value)
                value_cell.font = _HEADER_FONT
                value_cell.alignment = Alignment(horizontal="right", vertical="center")
                row_index += 1

            fixed_cols = max(len(headers), 2) if headers else 2
            for col in range(fixed_cols):
                col_rows = [r for r in rows if len(r) > col]
                if headers:
                    col_rows.append(headers)
                _auto_width(ws, col, col_rows)

            row_index += 2

        wb.save(file_path)
        return True
    except Exception as exc:  # noqa: BLE001
        logger.error(f"Failed to write Excel workbook '{file_path}': {exc}")
        return False